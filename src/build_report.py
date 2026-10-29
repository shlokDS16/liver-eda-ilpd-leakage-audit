"""Builds report/Liver_EDA_Report.docx from results/results.json and figures/.

Run after src/analysis.py:  python src/build_report.py
"""
import json
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from docx_theme import DISPLAY, INK, MUTED, OX, Report, _font

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "figures"
OUT = ROOT / "report"
OUT.mkdir(exist_ok=True)
R = json.load(open(ROOT / "results" / "results.json"))

M1, M2, M5, M6 = R["m1"], R["m2"], R["m5"], R["m6"]
A, B, C = R["m7"]["A"], R["m7"]["B"], R["m7"]["C"]
REG = R["m7_reg"]
POOL = R["m7_pool_rule"]


def pct(x, d=1):
    return f"{x * 100:.{d}f}%"


def n(x):
    return f"{int(x):,}"


def f2(x):
    return f"{x:.2f}"


def f3(x):
    return f"{x:.3f}"


SHORTN = {"Logistic regression": "Logistic regr.", "Gaussian naive Bayes": "Naive Bayes",
          "k-nearest neighbours (k=15)": "KNN (k = 15)", "Decision tree": "Decision tree",
          "Random forest": "Random forest", "XGBoost": "XGBoost", "SVM (RBF kernel)": "SVM (RBF)"}


def mrow(m, v):
    return [SHORTN[m], f"{v['cv_auc_mean']:.3f} ± {v['cv_auc_sd']:.3f}", f3(v["train_acc"]), f3(v["acc"]), f3(v["precision"]),
            f3(v["recall"]), f3(v["specificity"]), f3(v["f1"]), f3(v["bacc"]),
            f"{v['auc']:.3f} [{v['auc_ci'][0]:.2f}, {v['auc_ci'][1]:.2f}]"]


MHEAD = ["Model", "CV AUC", "Train acc", "Test acc", "Prec.", "Recall", "Spec.", "F1", "Bal. acc", "Test AUC [95% CI]"]
MW = [2.45, 2.35, 1.2, 1.2, 1.05, 1.1, 1.05, 1.0, 1.2, 2.8]
MNOTE = " Prec. = precision, Spec. = specificity, Bal. acc = balanced accuracy."

ORDER = ["f01_records", "f29_provenance", "f02_missing", "f03_imputation", "f06_skew", "f04_binning",
         "f05_distributions", "f07_boxplots", "f08_correlation", "f09_scatter", "f10_pairs",
         "f11_nd_heatmap", "f12_dotchart", "f13_bars",
         "f14_gmm_bic", "f15_dendrogram", "f16_clusters", "f17_outliers",
         "f18_pca", "f19_biplot", "f20_factors", "f21_embeddings", "f22_som",
         "f23_regression", "f25_nn_distance", "f24_cm_A", "f24_cm_B", "f24_cm_C", "f26_roc", "f27_protocol_gap",
         "f28_importance"]
F = {k: i + 1 for i, k in enumerate(ORDER)}


def rng_(a, b):
    return f"Figures {F[a]} to {F[b]}"


rep = Report()
d = rep.doc

# =========================================================================== cover
sec0 = d.sections[0]
sec0.top_margin = Cm(2.6)
rep.p("VELLORE INSTITUTE OF TECHNOLOGY", size=8, color=MUTED, font=DISPLAY, after=2)
rep.p("BCSE331L  EXPLORATORY DATA ANALYSIS  /  COURSE PROJECT", size=8, color=OX, font=DISPLAY, after=110)
t = rep.p("Thirty thousand rows,", size=34, color=INK, font=DISPLAY, after=0)
t.paragraph_format.line_spacing = 0.95
t = rep.p("five hundred and seventy patients", size=34, color=OX, font=DISPLAY, after=14)
t.paragraph_format.line_spacing = 0.95
rep.p("An exploratory analysis of the liver patient records, and why the model that scores 99.7 per cent "
      "learned almost nothing about the liver", size=13, italic=True, color=INK, after=28)
rule = rep.p("", after=18)
rep._rule(rule, "8C2F39", 8, 1)

kp = d.add_table(rows=2, cols=3)
nums = [(n(M1["n_rows"]), "rows in the course file"),
        (n(M2["orig_unique"]), "distinct patients the rows were built from"),
        (f"{C['models']['SVM (RBF kernel)']['cv_auc_mean']:.2f}", f"honest ROC-AUC, against {A['models']['Random forest']['auc']:.3f} on the file")]
for j, (big, small) in enumerate(nums):
    c0 = kp.rows[0].cells[j].paragraphs[0]
    r = c0.add_run(big)
    _font(r, DISPLAY, 26, OX if j == 2 else INK)
    c1 = kp.rows[1].cells[j].paragraphs[0]
    r = c1.add_run(small)
    _font(r, DISPLAY, 8.5, MUTED)
    for row in kp.rows:
        row.cells[j].width = Cm(5.4)
rep.p("", after=150)
for label, val in (("Submitted by", "Shlok Kumar Goenka"), ("Register number", "23BDS0316"),
                   ("Faculty", "Balaji G N"), ("School", "School of Computer Science and Engineering (SCOPE)"),
                   ("Submitted", "October 2026")):
    para = rep.p("", after=1)
    r = para.add_run(f"{label.upper()}   ")
    _font(r, DISPLAY, 7.5, MUTED)
    rep._tracking(r, 20)
    r = para.add_run(val)
    _font(r, DISPLAY, 10, INK)

# =========================================================================== front matter
rep.new_body_section("LIVER PATIENT RECORDS  /  EDA COURSE PROJECT  /  BCSE331L")

h = d.add_paragraph(style="Heading 1")
r = h.add_run("Abstract")
_font(r, DISPLAY, 20, INK)
rep._rule(h, "8C2F39", 8, 6)
pr = R["m7_pool_rule"]["A"]
rep.p(
    f"The course file holds {n(M1['n_rows'])} liver function records labelled as patient or non-patient. "
    f"I worked through it the way the syllabus is ordered: data types and scales first, then cleaning and missing "
    f"values, then summaries, plots, clustering, dimensionality reduction and, at the end, predictive models with a "
    f"proper train and test split. The cleaning stage turned up the most important fact in the project. "
    f"{n(M2['n_exact_dup'])} rows ({pct(M2['n_exact_dup'] / M1['n_rows'])}) are exact copies, and the "
    f"{n(M2['n_after_dedup'])} that remain were generated from the 583-record Indian Liver Patient Dataset "
    f"published on the UCI repository. Every lab value in an added row comes from an original patient of the "
    f"same class, so the values themselves give the label away.", align="justify")
rep.p(
    f"That explains the headline numbers. On a random 80/20 split of the file, random forest and XGBoost reach "
    f"{pct(A['models']['Random forest']['acc'])} test accuracy and an AUC of {f3(A['models']['Random forest']['auc'])}, "
    f"and a lookup rule with no learning in it reaches an AUC of {f2(pr['auc'])}. Trained and tested on the "
    f"{M2['orig_unique']} unique real patients, the same forest has a specificity of "
    f"{pct(C['models']['Random forest']['specificity'], 0)}. The best honest models are an RBF support vector "
    f"machine (cross-validated AUC {f3(C['models']['SVM (RBF kernel)']['cv_auc_mean'])}) and logistic regression "
    f"({f3(C['models']['Logistic regression']['cv_auc_mean'])}), which a McNemar test cannot separate "
    f"(p = {f2(C['mcnemar']['p'])}). I recommend logistic regression because it is as accurate and far easier to explain. "
    f"The exploratory work also shows that bilirubin, the transaminases and alkaline phosphatase carry most of the "
    f"signal, that albumin carries some in the opposite direction, and that the augmentation erased a real "
    f"gender and age effect present in the original patients.", align="justify")
kw = rep.p("", after=0)
r = kw.add_run("KEYWORDS   ")
_font(r, DISPLAY, 7.5, OX)
rep._tracking(r, 20)
r = kw.add_run("exploratory data analysis, liver function tests, missing data, EM algorithm, clustering, PCA, "
               "data leakage, synthetic data, model evaluation")
_font(r, DISPLAY, 9, MUTED)

rep.page_break()
h = d.add_paragraph()
h.paragraph_format.space_after = Pt(10)
r = h.add_run("Contents")
_font(r, DISPLAY, 20, INK)
rep._rule(h, "8C2F39", 8, 6)
rep.toc()

rep.page_break()
h = d.add_paragraph(style="Heading 2")
h.add_run("How the report follows the syllabus")
rep.p("Each chapter corresponds to one module of BCSE331L. The table lists the module topics and where each one is "
      "applied to this dataset. Two topics have no natural home in a single-snapshot tabular dataset (time series and "
      "document clustering); the chapters say so plainly instead of forcing an example.", align="justify")
rep.table(["Module", "Syllabus topics", "Where in this report"], [
    ["1  Introduction to EDA", "Steps in EDA, numerical and categorical data, measurement scales, EDA against classical and Bayesian analysis, tools", "Chapter 1, Tables 2 to 3"],
    ["2  Data transformation", "Deduplication, replacing values, discretisation and binning, missing data, traditional methods, maximum likelihood (EM)", f"Chapter 2, {rng_('f01_records', 'f04_binning')}"],
    ["3  Correlation analysis", "Univariate, bivariate and multivariate analysis; time series", f"Chapter 3, {rng_('f05_distributions', 'f10_pairs')}"],
    ["4  Summarisation and visualisation", "Summary measures, 1-D, 2-D and n-D analysis, contingency tables, scatter plots, dot charts, bar plots", f"Chapter 4, {rng_('f11_nd_heatmap', 'f13_bars')}"],
    ["5  Clustering", "Spectral, minimum spanning tree, model-based (EM), hierarchical agglomerative, outlier detection by clustering", f"Chapter 5, {rng_('f14_gmm_bic', 'f17_outliers')}"],
    ["6  Dimensionality reduction", "PCA, SVD, factor analysis, intrinsic dimensionality, MDS, manifold learning, self-organising maps", f"Chapter 6, {rng_('f18_pca', 'f22_som')}"],
    ["7  Model development and evaluation", "Linear regression and its evaluation, accuracy and what it hides, supervised against reinforcement learning", f"Chapter 7, {rng_('f23_regression', 'f28_importance')}"],
    ["8  Contemporary issues", "Synthetic data and leakage, fairness, reproducibility", "Chapter 8"],
], caption="Syllabus coverage map.", widths=[3.6, 8.6, 4.2], align=["left", "left", "left"], font_size=8)

# =========================================================================== 1
rep.h1("1", "The data and the questions", "Module 1  /  Introduction to EDA")
rep.p(
    "Liver disease in India is often found late, because the early stages have few symptoms and the routine blood "
    "panel that could flag it is read one number at a time. The dataset for this project is that panel: age, gender, "
    "two bilirubin measures, three enzymes, total protein, albumin and the albumin to globulin ratio, with a label "
    "saying whether the person was a liver patient. The file was downloaded from the course repository "
    "(salemprakash/EDA, Data/Liver Data.csv) on 14 October 2026.", align="justify")
rep.p("I set three questions at the start and kept them fixed:", after=3)
rep.bullets([
    "Is the file clean enough to trust, and what has to be done to it before any statistic means anything?",
    "Which measurements separate patients from non-patients, and do the patients fall into natural groups?",
    "How well can a model predict the label on people it has not seen, and which model should be used?",
], numbered=True)
rep.p(
    "EDA, as Tukey framed it and as the course textbook (Mukhiya and Ahmed, 2020) teaches it, runs in a loop: define "
    "the problem, prepare the data, summarise it, visualise it, model it, and go back when something looks wrong. "
    "This report went round that loop more than once. The first pass of modelling produced 99.7 per cent accuracy, "
    "which sent me back to the data preparation stage, and what I found there rewrote Chapter 7.", align="justify")

rep.h2("Variables, data types and measurement scales")
rep.table(["Variable", "Data type", "Scale", "Unit", "Note"], [
    ["Age", "Numerical, discrete", "Ratio", "years", "Ages above 89 were capped at 90 in the source"],
    ["Gender", "Categorical", "Nominal", "", "Male, Female, 902 blank"],
    ["Total bilirubin", "Numerical, continuous", "Ratio", "mg/dL", "Breakdown product of haemoglobin"],
    ["Direct bilirubin", "Numerical, continuous", "Ratio", "mg/dL", "Conjugated part of total bilirubin"],
    ["Alkaline phosphatase (ALP)", "Numerical, continuous", "Ratio", "IU/L", "Rises with bile duct obstruction"],
    ["ALT (SGPT)", "Numerical, continuous", "Ratio", "IU/L", "Enzyme released by damaged liver cells"],
    ["AST (SGOT)", "Numerical, continuous", "Ratio", "IU/L", "Also released by heart and muscle"],
    ["Total protein", "Numerical, continuous", "Ratio", "g/dL", "Albumin plus globulins"],
    ["Albumin", "Numerical, continuous", "Ratio", "g/dL", "Made by the liver; falls in chronic disease"],
    ["A/G ratio", "Numerical, continuous", "Ratio", "ratio", "Albumin divided by globulin"],
    ["Result (target)", "Categorical, binary", "Nominal", "", "1 = liver patient, 2 = not, recoded to 1 and 0"],
], caption="Variables in the file. The engineered bands in Chapter 2 (age band, bilirubin band) are ordinal.",
    widths=[3.6, 3.2, 1.5, 1.4, 6.7], align=["left"] * 5, font_size=8)
rep.p(
    "All nine lab measurements sit on a ratio scale, with a true zero, so ratios such as \"twice the normal ALT\" are "
    "meaningful and log transforms are allowed. Gender and the target are nominal. When I cut age and bilirubin into "
    "bands in Chapter 2 they become ordinal: the order of the bands matters but the gaps between them do not. "
    "Nothing in this dataset is on a pure interval scale; temperature in Celsius would be the classroom example.",
    align="justify")

rep.h2("Exploratory, classical and Bayesian analysis")
rep.p("The three approaches differ in the order of the steps, which is the clearest way to see them side by side.",
      after=3)
rep.table(["", "Classical (confirmatory)", "Exploratory", "Bayesian"], [
    ["Sequence", "Problem, data, model, analysis, conclusions", "Problem, data, analysis, model, conclusions", "Problem, data, model, prior, analysis, conclusions"],
    ["Model", "Chosen before seeing the data", "Suggested by the data", "Chosen beforehand, plus a prior"],
    ["Main output", "p-values, confidence intervals", "Plots, summaries, hypotheses", "Posterior distributions"],
    ["In this project", "Mann-Whitney, chi-square and McNemar tests", "Most of Chapters 2 to 6", "Gaussian naive Bayes in Chapter 7"],
], caption="Three ways to analyse the same data.", widths=[2.6, 4.6, 4.6, 4.6], align=["left"] * 4, font_size=8)

rep.h2("Tools")
rep.p(
    f"Everything runs in Python {R['env']['python']} with pandas {R['env']['pandas']}, NumPy {R['env']['numpy']}, "
    f"SciPy, scikit-learn {R['env']['sklearn']}, statsmodels {R['env']['statsmodels']}, XGBoost {R['env']['xgboost']} "
    f"and matplotlib. Two methods are written from scratch because no maintained library matched the syllabus "
    f"version: the EM algorithm for a multivariate normal with missing values, and the self-organising map. "
    f"The whole analysis is one script (src/analysis.py) that runs in about {int(R['env']['runtime_s'] // 60) + 1} "
    f"minutes with a fixed seed of {R['env']['seed']}.", align="justify")

# =========================================================================== 2
rep.h1("2", "Cleaning, missing values and binning", "Module 2  /  Data transformation")
rep.h2("Deduplication")
rep.p(
    f"The raw file has {n(M1['n_rows'])} rows and 11 columns. {n(M2['n_exact_dup'])} rows are exact copies of "
    f"another row, every column identical, which is {pct(M2['n_exact_dup'] / M1['n_rows'])} of the file. Removing "
    f"them leaves {n(M2['n_after_dedup'])} rows. The class balance barely moves (71.4 per cent patients before, "
    f"{pct(M2['result_counts_dedup']['1'] / M2['n_after_dedup'])} after), so the copies were spread evenly. "
    f"Duplicates are not harmless here. With a random split, a copy of a test row will often sit in the training "
    f"set, and the model is then graded on a question it has already seen the answer to.", align="justify")
rep.figure(FIG / "f01_records.png", "Records before and after removing exact duplicates, next to the original UCI "
           "release. The class mix is almost identical in all three.", 15.5)

rep.h2("Where the extra rows came from")
pv = M2["provenance"]
rep.p(
    f"Even after deduplication the file is 33 times larger than any liver dataset I could find with these columns. "
    f"The first rows matched the Indian Liver Patient Dataset (Ramana, Babu and Venkateswarlu, 2012), which has 583 "
    f"records, so I downloaded that release and compared. {M2['n_orig_rows_found']} of the original rows appear "
    f"verbatim in the course file. The other {n(M2['n_synthetic'])} are new combinations, and they follow a strict "
    f"pattern: every bilirubin, enzyme and protein value in an added row also appears in an original patient of the "
    f"same class ({pct(pv['alkphos']['same_class'], 0)} for every lab column). Far fewer appear in the other class; for "
    f"alkaline phosphatase only {pct(pv['alkphos']['other_class'], 0)} do (Figure {F['f29_provenance']}).", align="justify")
rep.p(
    "That is what you get if each column of a synthetic row is copied from a random patient of the chosen class. "
    "SMOTE, the usual oversampling method, would interpolate and produce values that never occurred, so this is a "
    "different, simpler generator. It has two side effects that drive the rest of the report. The label leaks "
    "through the values, because a value seen only among non-patients can never appear in a patient row. And any "
    "relationship between columns that the generator ignored is diluted towards zero.", align="justify")
rep.figure(FIG / "f29_provenance.png", "Left: share of values in the added rows that also occur among original "
           "patients of the same class (red) and of the other class (grey). Right: the original data has a gender gap "
           "in disease rate; the 30k file does not.", 16.4)
rep.note("Finding",
         f"In the original 583 patients, men have a disease rate of {pct(M2['orig_gender_rate']['Male'])} and women "
         f"{pct(M2['orig_gender_rate']['Female'])}, and age correlates with the label (r = {f2(M2['orig_age_corr'])}, "
         f"Mann-Whitney p = {M2['orig_age_p']:.3f}). In the course file both effects are gone: "
         f"{pct(M2['aug_gender_rate']['Male'])} against {pct(M2['aug_gender_rate']['Female'])}, and r = "
         f"{f3(M2['aug_age_corr'])}. Any EDA conclusion about age or gender drawn from the 30k file alone would be wrong.")

rep.h2("Replacing values and consistency checks")
rep.p(
    f"The target was coded 1 for a patient and 2 for a non-patient. I recoded it to 1 and 0 so that means and "
    f"model probabilities read as disease rates. Column names had stray leading spaces (\" Alkphos Alkaline "
    f"Phosphotase\") and a misspelling (\"Protiens\"), so every column got a short name. Two clinical rules make "
    f"useful sanity checks. Direct bilirubin is part of total bilirubin, so it cannot be larger: "
    f"{M2['db_gt_tb']} rows break this rule, probably rounding in the source. The A/G ratio should equal "
    f"albumin divided by (total protein minus albumin), and it does within 0.1 for "
    f"{pct(M2['ag_consistent_share'])} of the {n(M2['ag_checked'])} rows where all three are present. I kept the "
    f"rows that break these rules because the error is small and the rules are only approximate at one decimal place.",
    align="justify")

rep.h2("Missing data")
ms = M2["missing"]
rep.p(
    f"After deduplication {n(M2['cells_missing'])} cells are blank, spread over {n(M2['rows_any_missing'])} rows "
    f"({pct(M2['rows_any_missing'] / M2['n_after_dedup'])}). No column is missing more than "
    f"{pct(max(v['share'] for v in ms.values()))} and no row is missing more than {M2['max_missing_per_row']} of its "
    f"ten values. Gender is the most often missing ({n(ms['gender']['missing'])} rows), age the least (2 rows).",
    align="justify")
rows = []
for c, lab in (("gender", "Gender"), ("alkphos", "Alkaline phosphatase"), ("tb", "Total bilirubin"), ("db", "Direct bilirubin"),
               ("sgpt", "ALT (SGPT)"), ("ag", "A/G ratio"), ("alb", "Albumin"), ("sgot", "AST (SGOT)"),
               ("tp", "Total protein"), ("age", "Age")):
    v = ms[c]
    rows.append([lab, n(v["missing"]), pct(v["share"], 2), pct(v["disease_rate_if_missing"]),
                 pct(v["disease_rate_if_observed"]), f"{v['chi2_p']:.3f}"])
rep.table(["Variable", "Missing", "Share", "Disease rate, missing", "Disease rate, observed", "Chi-square p"], rows,
          caption="Missing values per column, and whether rows with a gap have a different disease rate.",
          widths=[3.8, 1.8, 1.8, 3.0, 3.0, 2.2], font_size=8)
rep.p(
    "Rubin (1976) separates three mechanisms. Missing completely at random (MCAR) means the gap has nothing to do "
    "with any value; missing at random (MAR) means it depends on other observed columns; missing not at random "
    "(MNAR) means it depends on the hidden value itself. The last two columns of Table 4 test one consequence of "
    "MCAR: if gaps were pure chance, rows with a gap would have the same disease rate as rows without. For most "
    "columns they do. Rows missing ALT are "
    f"{abs(ms['sgpt']['disease_rate_if_missing'] - ms['sgpt']['disease_rate_if_observed']) * 100:.0f} points less "
    f"likely to be patients (chi-square p = {ms['sgpt']['chi2_p']:.3f}), and the two bilirubins and the A/G ratio show "
    "gaps of about 4 points with p between 0.05 and 0.07. That is evidence against MCAR for at least one column, so "
    "simply deleting incomplete rows would tilt the class mix slightly. An imputation that uses the other columns is the "
    "safer choice.", align="justify")
rep.figure(FIG / "f02_missing.png", "Left: share of each column that is blank. Right: disease rate among rows where the "
           "value is present (teal) and absent (red). Age has only two blank rows, so its 100 per cent means nothing.", 16.4)

rep.h3("Comparing imputation methods")
ir = M2["imputation_rmse"]
rep.p(
    f"To compare methods fairly I took the {n(M2['n_after_dedup'] - M2['rows_any_missing'])} complete rows, hid "
    f"{n(M2['imputation_cells_hidden'])} values at random (10 per cent), filled them with each method and measured "
    f"the error against the truth. Enzymes and bilirubin were logged and every column standardised first, so the "
    f"error is in standard deviations. Mean and median imputation are the traditional methods in the syllabus. "
    f"The maximum likelihood method is the EM algorithm of Dempster, Laird and Rubin (1977) for a multivariate "
    f"normal: the E step fills each gap with its conditional expectation given the observed values in that row, "
    f"the M step re-estimates the mean vector and covariance matrix, including the extra variance from the filled "
    f"cells, and the two repeat until the estimates stop changing.", align="justify")
rep.figure(FIG / "f03_imputation.png", "Error on deliberately hidden values. Lower is better; one unit is one "
           "standard deviation of the column.", 15.5)
rep.p(
    f"Mean ({f2(ir['Mean'])}) and median ({f2(ir['Median'])}) are no better than guessing the average, which is "
    f"what they are. EM roughly halves the error ({f2(ir['EM, multivariate normal (MLE)'])}) because it uses the "
    f"strong correlations between the bilirubins and between the enzymes; chained regression lands in the same "
    f"place ({f2(ir['Iterative (chained regression)'])}). KNN does best ({f2(ir['KNN (k=5)'])}), and after the "
    f"provenance finding the reason is plain: a row's nearest neighbours are often siblings built from the same "
    f"original patients, so KNN is partly copying the answer. I used EM for the descriptive work because its "
    f"assumptions are explicit and it converged in {M2['em_iterations']} iterations, and filled gender with the "
    f"mode ({M2['gender_mode']}). Inside the models in Chapter 7 the imputer is fitted on the training rows only.",
    align="justify")

rep.h2("Log transformation")
sk = R["m3_skew"]
rep.p(
    f"The five enzyme and bilirubin columns are heavily right-skewed. AST has a skewness of {f2(sk['sgot']['before'])}; "
    f"its median is {R['m4_summary']['sgot']['median']:g} IU/L and its maximum is {n(R['m4_summary']['sgot']['max'])}. A log(1 + x) transform brings "
    f"all five to 1.7 or less (Figure {F['f06_skew']}). The protein columns and age are close to symmetric and were left alone. All "
    f"correlations, clustering and dimensionality reduction from here on use the logged values.", align="justify")
rep.figure(FIG / "f06_skew.png", "Skewness of each column before (red) and after (teal) the log transform.", 13.0)

rep.h2("Discretisation and binning")
bins = M2["bins"]
rep.p(
    "Binning trades detail for readability. I made three kinds of bins. Age was cut at clinical boundaries (under 18, "
    "18 to 39, 40 to 59, 60 and over). Total bilirubin was cut at 1.2 mg/dL, the usual upper reference limit, and "
    "at 3 mg/dL, roughly where jaundice becomes visible. Albumin was cut both into four equal-width bins and into "
    "four equal-frequency bins (quartiles) to show the difference.", align="justify")
tb = bins["tb_band"]
ew = bins["alb_eqwidth"]
rep.p(
    f"The bilirubin bands separate the classes cleanly: the disease rate climbs from "
    f"{pct(tb['Normal (1.2 or less)']['mean'], 0)} to {pct(tb['Raised (1.2 to 3)']['mean'], 0)} to "
    f"{pct(tb['High (over 3)']['mean'], 0)}. Age bands do nothing, which matches the provenance finding. Equal-width "
    f"bins on albumin put only {min(int(v['size']) for v in ew.values())} rows in the sparsest bin and "
    f"{n(max(int(v['size']) for v in ew.values()))} in the fullest, while the quartiles hold between "
    f"{n(min(int(v['size']) for v in bins['alb_eqfreq'].values()))} and "
    f"{n(max(int(v['size']) for v in bins['alb_eqfreq'].values()))}, so "
    f"quartiles give the steadier rates when a variable is bunched in the middle.", align="justify")
rep.figure(FIG / "f04_binning.png", "Disease rate by age band, bilirubin band and albumin quartile.", 16.4)

# =========================================================================== 3
rep.h1("3", "One variable, two variables, many", "Module 3  /  Correlation analysis")
rep.h2("Univariate analysis")
rep.p(
    f"Figure {F['f05_distributions']} draws each measurement separately for patients and non-patients, with the enzymes and bilirubins on a "
    "log axis. Three things stand out. The patient curves for bilirubin, ALP, ALT and AST have long right tails that "
    "the non-patient curves lack. The non-patient curves are narrow, which suggests most of them sit inside reference "
    "ranges. And age, total protein and albumin overlap almost completely.", align="justify")
rep.figure(FIG / "f05_distributions.png", "Distribution of each variable by class. Red: liver patients. Teal: "
           "non-patients. Five panels use a log axis.", 16.4)
iq = R["m3_iqr_outliers"]
rep.p(
    f"The 1.5 IQR rule flags {n(iq['tb']['raw_scale'])} total bilirubin values as outliers on the raw scale but only "
    f"{n(iq['tb']['log_scale'])} on the log scale; for AST the drop is from {n(iq['sgot']['raw_scale'])} to "
    f"{n(iq['sgot']['log_scale'])}. On a skewed variable the rule mostly flags the tail of the distribution, not "
    f"errors. A bilirubin of 75 mg/dL is extreme but clinically possible in obstructive jaundice, so I treated these "
    f"as real and did not delete them. Chapter 5 revisits outliers with clustering.", align="justify")

rep.h2("Bivariate analysis against the target")
mw = R["m3_mannwhitney"]
rows = []
for k, v in sorted(mw.items(), key=lambda kv: -abs(kv[1]["rank_biserial"])):
    p = "< 0.001" if v["p"] < 0.001 else f"{v['p']:.2f}"
    rows.append([k, f"{v['median_disease']:g}", f"{v['median_no_disease']:g}", p, f"{v['auc']:.3f}", f"{v['rank_biserial']:+.3f}"])
rep.table(["Variable", "Median, patients", "Median, others", "Mann-Whitney p", "AUC", "Rank-biserial r"], rows,
          caption="Each variable compared between the classes. AUC is the chance that a random patient has a higher "
                  "value than a random non-patient.", widths=[3.8, 2.6, 2.6, 2.6, 1.8, 2.6], font_size=8)
rep.p(
    "Because the distributions are skewed I used the Mann-Whitney U test rather than a t-test. With 19,000 rows almost "
    "anything is \"significant\", so the effect size matters more than the p-value. Total bilirubin and AST lead (rank-"
    f"biserial r about 0.40), followed by direct bilirubin, ALT and ALP. Albumin and the A/G ratio run the other way: "
    f"patients have less albumin (median {mw['Albumin']['median_disease']:g} against "
    f"{mw['Albumin']['median_no_disease']:g} g/dL), which fits a liver that is making less protein. Total protein "
    f"is statistically significant and practically useless (r = {mw['Total protein']['rank_biserial']:+.3f}). Age has "
    f"no effect at all in this file.", align="justify")
rep.figure(FIG / "f07_boxplots.png", "Box plots by class. D: liver patient. N: non-patient. Enzyme and bilirubin "
           "panels use a log axis.", 16.4)

rep.h2("Correlation between the measurements")
cr = R["m3_corr"]
rep.p(
    f"Pearson correlation assumes a straight-line relationship and is pulled around by extreme values, so I computed "
    f"it on the logged values and set Spearman's rank correlation beside it (Figure {F['f08_correlation']}). They agree closely. Total "
    f"and direct bilirubin are almost the same measurement (r = {f2(cr['pearson_log']['tb']['db'])} on the log scale), "
    f"ALT and AST move together (r = {f2(cr['pearson_log']['sgpt']['sgot'])}), and albumin, total protein and A/G "
    f"ratio form a third block. Across blocks the correlations are weak. On the raw scale the same two pairs give "
    f"{f2(cr['pearson_raw_tb_db'])} and {f2(cr['pearson_raw_sgpt_sgot'])}: a few extreme readings pull the raw "
    f"coefficients down, which is one more reason to log these columns first.", align="justify")
rep.figure(FIG / "f08_correlation.png", "Correlation matrices. Left: Pearson on logged enzymes and bilirubins. "
           "Right: Spearman rank correlation on raw values.", 16.4)
rep.figure(FIG / "f09_scatter.png", "Scatter plots of the three strongest pairs, 4,000 rows sampled. Patients "
           "stretch far along both bilirubin axes; non-patients stay in the lower-left corner.", 16.4)

rep.h2("Multivariate view")
rep.p(
    f"A pair plot (Figure {F['f10_pairs']}) shows every two-way combination of four variables at once. The useful reading is the "
    "bottom-left of each panel: non-patients cluster where bilirubin and AST are low and albumin is normal, and "
    "patients spread out from there in every direction. No single pair splits the classes. That already hints that a "
    "linear model will find a modest boundary, and that the classes overlap heavily in the middle.", align="justify")
rep.figure(FIG / "f10_pairs.png", "Pair plot of total bilirubin, AST, albumin and age (1,500 rows sampled).", 14.5)

rep.h2("A note on time series")
rep.p(
    "Module 3 also covers time series: time-based indexing, resampling and grouping by period. This dataset has no "
    "time column. Each row is one patient at one visit, with no date of collection and no repeated measurements, so "
    "there is nothing to index by time or resample. Sorting by age would not create a time series, because the rows "
    "are different people. A natural extension would be a cohort with repeated liver panels per patient, where "
    "resampling to monthly means and plotting ALT trajectories would apply directly.", align="justify")

# =========================================================================== 4
rep.h1("4", "Summaries, tables and charts", "Module 4  /  Data summarisation and visualisation")
rep.h2("Statistical summary measures")
sm_ = R["m4_summary"]
rows = []
names = {"age": "Age", "tb": "Total bilirubin", "db": "Direct bilirubin", "alkphos": "ALP", "sgpt": "ALT (SGPT)",
         "sgot": "AST (SGOT)", "tp": "Total protein", "alb": "Albumin", "ag": "A/G ratio"}
for c, lab in names.items():
    v = sm_[c]
    rows.append([lab, n(v["n"]), f"{v['mean']:.2f}", f"{v['sd']:.2f}", f"{v['median']:g}", f"{v['q1']:g} to {v['q3']:g}",
                 f"{v['min']:g} to {v['max']:g}", f"{v['cv']:.2f}", f"{v['skew']:.2f}", f"{v['kurtosis']:.1f}"])
rep.table(["Variable", "n", "Mean", "SD", "Median", "IQR", "Range", "CV", "Skew", "Kurtosis"], rows,
          caption="Summary of the observed values after deduplication (no imputation).",
          widths=[2.8, 1.4, 1.3, 1.3, 1.3, 2.1, 2.3, 1.1, 1.2, 1.6], font_size=7.8)
rep.p(
    f"Mean and median tell different stories for the skewed columns. Mean AST is {sm_['sgot']['mean']:.0f} IU/L "
    f"while the median is {sm_['sgot']['median']:g}; a few hundred very high readings drag the mean up. The "
    f"coefficient of variation (SD over mean) puts the spread on a common footing: "
    f"{min(sm_[c]['cv'] for c in ('tb', 'db', 'sgpt', 'sgot')):.1f} or more for the bilirubins, ALT and AST, "
    f"{sm_['alkphos']['cv']:.2f} for ALP, and {max(sm_[c]['cv'] for c in ('age', 'tp', 'alb', 'ag')):.2f} or less "
    f"for age and the protein columns. For reporting these columns the median and IQR are the honest "
    f"pair.", align="justify")

rep.h2("Two-way contingency tables")
ct = R["m4_contingency"]
g = ct["gender"]["table"]
t_ = ct["tb_band"]["table"]
rows = [["Female", n(g["Female"]["No disease"]), n(g["Female"]["Liver disease"]), pct(g["Female"]["Disease rate"])],
        ["Male", n(g["Male"]["No disease"]), n(g["Male"]["Liver disease"]), pct(g["Male"]["Disease rate"])]]
for k in ("Normal (1.2 or less)", "Raised (1.2 to 3)", "High (over 3)"):
    rows.append([f"Bilirubin {k[0].lower()}{k[1:]}", n(t_[k]["No disease"]), n(t_[k]["Liver disease"]), pct(t_[k]["Disease rate"])])
rep.table(["Group", "Non-patients", "Patients", "Disease rate"], rows,
          caption="Contingency tables of gender and bilirubin band against the target.",
          widths=[5.6, 3.0, 3.0, 3.0], font_size=8.5)
rep.p(
    f"For gender the chi-square test finds nothing (p = {ct['gender']['p']:.2f}, Cramér's V = "
    f"{ct['gender']['cramers_v']:.3f}), and the result holds when the mode-imputed rows are dropped (p = "
    f"{ct['gender_observed_only_p']:.2f}). The bilirubin band is strongly associated (chi-square = "
    f"{ct['tb_band']['chi2']:,.0f} on {ct['tb_band']['dof']} degrees of freedom, V = {ct['tb_band']['cramers_v']:.2f}). "
    f"Cramér's V is the more useful number here, since with 19,000 rows any chi-square test has enough power to "
    f"reject.", align="justify")

rep.h2("n-dimensional view")
rep.p(
    "Crossing three factors at once (age band, gender and bilirubin band) gives 24 cells, each with its own disease "
    f"rate. Figure {F['f11_nd_heatmap']} shows the result as a heat map. Reading across a row, the rate climbs with bilirubin in every "
    "age and gender group, from about 60 per cent to about 96 per cent. Reading down a column, almost nothing "
    "changes. The heat map makes the main structure of the file obvious in one glance: bilirubin matters, and age and "
    "gender do not, at least in the augmented data.", align="justify")
rep.figure(FIG / "f11_nd_heatmap.png", "Disease rate in each age, gender and bilirubin cell, with cell counts.", 14.0)

rep.h2("Dot chart and bar plot")
rep.p(
    "A dot chart (Cleveland's design) carries the same information as a bar chart with less ink, and it suits values "
    f"that can be negative. Figure {F['f12_dotchart']} ranks the nine variables by rank-biserial correlation with the target. The bar "
    f"plot in Figure {F['f13_bars']} shows disease rate by age band and gender with 95 per cent Wilson intervals. The intervals for "
    "the under-18 groups are wide because there are only a few hundred children.", align="justify")
rep.figure(FIG / "f12_dotchart.png", "Dot chart of effect sizes. Positive values are higher in patients.", 13.5)
rep.figure(FIG / "f13_bars.png", "Disease rate by age band and gender with 95 per cent Wilson intervals.", 13.5)

# =========================================================================== 5
rep.h1("5", "Looking for groups without the label", "Module 5  /  Clustering")
rep.p(
    "Clustering ignores the label and asks whether the records fall into groups on their own. All methods here work on "
    "the nine logged and standardised lab values. Spectral, hierarchical and minimum spanning tree clustering need a "
    "full distance matrix, so they ran on a random sample of 2,000 rows; K-means and the Gaussian mixture ran on "
    "everything. Document clustering, also in Module 5, needs text, and this dataset has none.", align="justify")
rep.h2("Model-based clustering with EM")
rep.p(
    f"A Gaussian mixture model treats the data as a blend of several multivariate normal groups and fits them with "
    f"the Expectation-Maximisation algorithm: the E step works out how likely each row is to belong to each group, the "
    f"M step re-estimates each group's mean and covariance from those weights. The Bayesian information criterion "
    f"(BIC) trades fit against the number of parameters. It kept falling up to {M5['k_bic']} components, the largest "
    f"I tried (Figure {F['f14_gmm_bic']}). With 19,000 rows BIC can afford many components, and part of what it is fitting is the "
    f"lumpy structure the generator created. The large drops end after three or four components, so I used 4 groups for every method to keep "
    f"them comparable.", align="justify")
rep.figure(FIG / "f14_gmm_bic.png", "BIC against the number of mixture components. Lower is better.", 12.0)
gp = M5["gmm_profile"]
rows = []
for i_, (k, v) in enumerate(gp.items()):
    rows.append([f"Group {i_ + 1}", n(v["n"]), pct(v["disease_rate"]), f"{v['tb']:g}", f"{v['sgpt']:g}", f"{v['sgot']:g}",
                 f"{v['alkphos']:g}", f"{v['alb']:g}"])
rep.table(["Group", "Rows", "Disease rate", "TBil", "ALT", "AST", "ALP", "Albumin"], rows,
          caption="The four Gaussian mixture clusters, numbered from highest to lowest disease rate (medians of each "
                  "variable).",
          widths=[2.6, 1.8, 2.2, 1.6, 1.6, 1.6, 1.8, 2.0], font_size=8.5)
rep.p(
    f"The clusters read like clinical stages. The largest group ({n(max(v['n'] for v in gp.values()))} rows) has "
    f"normal bilirubin and enzymes and the lowest disease rate. The highest-risk group has every enzyme raised and "
    f"bilirubin around 3.5 mg/dL, and nine in ten of its members are patients. A smaller group of about 1,000 "
    f"rows has the highest bilirubin of all (median 5.0 mg/dL) and the lowest albumin (2.6 g/dL) but only mildly "
    f"raised ALT, a pattern that fits chronic rather than acute damage.",
    align="justify")

rep.h2("Hierarchical, spectral and minimum spanning tree clustering")
rep.p(
    "Hierarchical agglomerative clustering starts with every row alone and merges the closest pair repeatedly. With "
    "Ward's criterion each merge is the one that adds least to the within-group variance, which is the hierarchical "
    f"version of a spherical Gaussian mixture (Fraley, 1998). The dendrogram in Figure {F['f15_dendrogram']} "
    "shows the last 30 merges, with a dashed line at the four-cluster cut. Spectral clustering builds a 15-nearest-neighbour "
    "graph, takes the leading eigenvectors of its Laplacian and clusters in that space, which lets it follow curved "
    "groups. Minimum spanning tree (MST) clustering connects all points with the shortest possible tree and cuts "
    "the longest edges (Zahn, 1971).", align="justify")
rep.figure(FIG / "f15_dendrogram.png", "Ward dendrogram, truncated to the last 30 merges. The dashed line is the "
           "four-cluster cut.", 16.0)
cl = M5["methods"]
rows = []
for k, v in cl.items():
    rows.append([k, v["sizes"].replace("/", " / "), f"{v['silhouette']:.3f}", f"{v['ari_vs_outcome']:+.3f}",
                 f"{pct(v['min_cluster_disease_rate'], 0)} to {pct(v['max_cluster_disease_rate'], 0)}"])
rep.table(["Method", "Cluster sizes (2,000-row sample)", "Silhouette", "ARI with label", "Disease rate range"], rows,
          caption="Five clustering methods compared on the same sample with four clusters.",
          widths=[4.0, 4.6, 2.0, 2.4, 3.4], font_size=8.2)
rep.p(
    f"No method recovers the label: the adjusted Rand index against it never exceeds "
    f"{max(v['ari_vs_outcome'] for v in cl.values()):.2f}. That is expected, because patients are not one group; "
    f"they are everyone whose numbers have drifted away from normal in any direction. The clusters are still useful, "
    f"since the disease rate ranges from about 50 to 96 per cent across them. MST clustering behaves differently from "
    f"the rest. Its longest edges ({f2(M5['mst_longest_edges'][-1])} against a median edge of "
    f"{f2(M5['mst_median_edge'])}) lead to small outlying clumps, so cutting them peels off groups of "
    f"{M5['methods']['Minimum spanning tree']['sizes'].split('/', 1)[1].replace('/', ', ')} records and leaves one "
    f"giant cluster. That makes it a poor partitioning method on this data and a good outlier detector.",
    align="justify")
rep.figure(FIG / "f16_clusters.png", "The same 2,000 rows projected onto the first two principal components and "
           "coloured by each method's clusters. The last panel shows the true label for comparison.", 16.4)

rep.h2("Outlier detection using clustering")
ol = M5["outliers"]
rep.p(
    f"Three clustering-based rules flagged unusual records: the 1 per cent of rows with the lowest density under the "
    f"Gaussian mixture ({ol['gmm_low_density']} rows), the 1 per cent farthest from their K-means centre "
    f"({ol['kmeans_far']}), and DBSCAN's noise points ({ol['dbscan_noise']}). The two density-based ideas agree only "
    f"partly (Jaccard {f2(ol['jaccard_gmm_km'])} between GMM and K-means). The flagged rows are not data errors. "
    f"{pct(ol['disease_rate_gmm_outliers'], 0)} of the GMM outliers are patients, with median bilirubin "
    f"{ol['median_tb_outliers']:.1f} mg/dL and median AST {ol['median_sgot_outliers']:.0f} IU/L. They are the "
    f"sickest people in the file. Removing them as \"noise\" would delete the clearest cases of the disease.",
    align="justify")
rep.figure(FIG / "f17_outliers.png", "Records flagged by low mixture density and by distance from the K-means "
           "centre, on the first two principal components.", 13.5)

# =========================================================================== 6
rep.h1("6", "Fewer dimensions", "Module 6  /  Dimensionality reduction")
rep.h2("Principal component analysis and SVD")
pv_ = M6["pca_var"]
rep.p(
    f"PCA rotates the nine standardised measurements into uncorrelated components ordered by variance. The first "
    f"component explains {pct(pv_[0])} of the variance, the second {pct(pv_[1])}. {M6['kaiser']} components have an "
    f"eigenvalue above 1 (Kaiser's rule), {M6['n90']} are needed for 90 per cent and {M6['n95']} for 95 per cent. "
    f"PC1 loads positively on both bilirubins and all three enzymes and negatively on albumin and A/G ratio, so it is "
    f"a general \"liver damage\" axis. PC2 is the protein axis. PC3 is almost entirely age (loading "
    f"{M6['loadings']['PC3']['Age']:.2f}), which tells you age shares nothing with the lab values.", align="justify")
rep.p(
    f"PCA and singular value decomposition are the same computation seen from two sides. If the centred data matrix "
    f"is X = U Σ Vᵀ, the columns of V are the principal directions and the squared singular values divided by n - 1 "
    f"are the eigenvalues. I checked this numerically: the largest difference between the two sets of eigenvalues was "
    f"{M6['svd_eig_maxdiff']:.1e}, which is rounding error.", align="justify")
rep.figure(FIG / "f18_pca.png", "Left: variance explained per component (bars) and cumulative (line). Right: "
           "loadings of the first four components.", 16.4)
rep.figure(FIG / "f19_biplot.png", "Biplot of the first two components. Arrows show how each variable loads; "
           "points are 2,000 sampled rows.", 13.5)

rep.h2("Factor analysis")
fa = M6["fa"]
rep.p(
    f"Factor analysis assumes each measurement is a mix of a few hidden factors plus its own noise, and it reports "
    f"that noise as the uniqueness. With three factors and varimax rotation the structure is clean (Figure {F['f20_factors']}). "
    f"Factor 1 is protein synthesis: albumin {fa['Factor 1']['Albumin']:+.2f}, total protein "
    f"{fa['Factor 1']['Total protein']:+.2f}, A/G ratio {fa['Factor 1']['A/G ratio']:+.2f}. Factor 2 is "
    f"hepatocellular injury: ALT {fa['Factor 2']['ALT (SGPT)']:+.2f}, AST {fa['Factor 2']['AST (SGOT)']:+.2f}. "
    f"Factor 3 is bilirubin handling: total {fa['Factor 3']['Total bilirubin']:+.2f}, direct "
    f"{fa['Factor 3']['Direct bilirubin']:+.2f}. These three are the textbook groupings of the liver panel "
    f"(synthesis, injury, excretion), recovered from data alone. Age has a uniqueness of "
    f"{fa['Uniqueness']['Age']:.2f}, so none of it is shared, and ALP sits mostly on its own "
    f"({fa['Uniqueness']['Alkaline phosphatase']:.2f}).", align="justify")
rep.figure(FIG / "f20_factors.png", "Varimax-rotated loadings of a three-factor model.", 11.0)

rep.h2("Intrinsic dimensionality")
mi = M6["mle_id"]
rep.p(
    f"Nine columns does not mean nine dimensions of real variation. The maximum likelihood estimator of Levina and "
    f"Bickel (2004) looks at how fast the number of neighbours grows with distance around each point. On 4,000 "
    f"unique rows it gives {mi['10']:.1f}, {mi['15']:.1f} and {mi['20']:.1f} for 10, 15 and 20 neighbours. "
    f"PCA's linear count was {M6['n90']} components for 90 per cent of the variance. The two answers fit together: "
    f"three groups of tightly linked measurements (proteins, enzymes, bilirubins) plus age and ALP collapse nine "
    f"columns to about four effective directions.", align="justify")

rep.h2("Non-linear methods: MDS, Isomap and t-SNE")
ek = M6["embed_knn_acc"]
rep.p(
    "Metric multidimensional scaling places points in two dimensions so that their distances match the original "
    "distances as closely as possible. Isomap does the same with distances measured along a neighbourhood graph, so it "
    "can unroll curved structure (Tenenbaum, de Silva and Langford, 2000). t-SNE keeps local neighbourhoods and lets "
    "global distances go (van der Maaten and Hinton, 2008). To compare them I trained a 5-nearest-neighbour classifier "
    "inside each two-dimensional map and measured how well it recovered the label.", align="justify")
rep.figure(FIG / "f21_embeddings.png", "Three non-linear two-dimensional maps of the same records, coloured by "
           "class.", 16.4)
rep.p(
    f"The plain linear projection scored best ({pct(ek['PCA (2 components)'], 0)}), ahead of t-SNE "
    f"({pct(ek['t-SNE (perplexity 30)'], 0)}), Isomap ({pct(ek['Isomap (12 neighbours)'], 0)}) and MDS "
    f"({pct(ek['Metric MDS'], 0)}, on 700 rows). Non-linear maps are not automatically better. Here the main "
    f"direction of difference between the classes is the straight \"damage\" axis that PC1 already captures. "
    f"t-SNE breaks the data into many small islands, which fits the family structure of the synthetic rows, but "
    f"the islands do not line up with the label.", align="justify")

rep.h2("Self-organising map")
so = M6["som"]
rep.p(
    f"A self-organising map (Kohonen, 1982) is a small neural network laid out as a 10 by 10 grid. Each node holds a "
    f"prototype vector; for each training record the closest node and its grid neighbours move a little towards it, "
    f"with the step size and neighbourhood shrinking over six passes. Similar records end up on nearby nodes. The "
    f"final quantisation error was {f2(so['quantisation_error'])} standard units and no node was empty. The U-matrix "
    f"in Figure {F['f22_som']} marks boundaries between regions in dark grey. Colouring nodes by disease rate shows a region of "
    f"{so['nodes_rate_below_50']} nodes where non-patients are the majority (lowest rate "
    f"{pct(so['node_rate_min'], 0)}), and nodes elsewhere where every record mapped is a patient.", align="justify")
rep.figure(FIG / "f22_som.png", "Self-organising map. Left: U-matrix (average distance from each node to its "
           "neighbours). Right: disease rate of the records mapped to each node.", 14.5)

# =========================================================================== 7
rep.h1("7", "Models and how to judge them", "Module 7  /  Model development and evaluation")
rep.h2("A linear regression model")
rc = REG["coef"]
rep.p(
    f"Before classification, Module 7 asks for a linear regression built and evaluated properly. I predicted log AST "
    f"from log ALT, log ALP, log total bilirubin, age, albumin and gender, on the {n(REG['n_train'] + REG['n_test'])} "
    f"complete rows split 80/20 at random. Logs make sense twice over: the enzymes are skewed, and a coefficient on "
    f"a log-log pair reads as an elasticity.", align="justify")
rows = []
lab = {"const": "Intercept", "log_sgpt": "log ALT", "log_alkphos": "log ALP", "log_tb": "log total bilirubin",
       "age": "Age (years)", "alb": "Albumin (g/dL)", "male": "Male"}
for k, v in rc.items():
    p = "< 0.001" if v["p"] < 0.001 else f"{v['p']:.2f}"
    rows.append([lab[k], f"{v['b']:+.4f}", f"{v['se']:.4f}", f"{v['t']:.1f}", p])
rep.table(["Term", "Coefficient", "Std. error", "t", "p"], rows,
          caption="Ordinary least squares fit for log AST on the training rows.", widths=[4.6, 2.8, 2.6, 2.0, 2.0],
          font_size=8.5)
rep.p(
    f"The model explains {pct(REG['r2_train'], 1)} of the variance in training and {pct(REG['r2_test'], 1)} on the "
    f"held-out rows, so it does not overfit. Almost all of that comes from ALT alone (R² = "
    f"{f2(REG['simple_r2_sgpt_only'])}). The ALT coefficient of {rc['log_sgpt']['b']:.2f} means a 10 per cent rise "
    f"in ALT goes with roughly an {rc['log_sgpt']['b'] * 10:.0f} per cent rise in AST. Age and gender are not "
    f"significant. On the original scale the typical error is large: the median absolute percentage error is "
    f"{pct(REG['median_ape_test'], 0)} and only {pct(REG['within_25pct'], 0)} of predictions land within 25 per cent "
    f"of the measured AST. A high R² on a log scale can hide a lot of error in the units a doctor reads. The residual "
    f"plot is centred on zero with roughly even spread, so the linear form is adequate.", align="justify")
rep.figure(FIG / "f23_regression.png", "Linear regression on the test rows. Left: predicted against measured AST "
           "(both axes logarithmic). Right: residuals.", 15.5)

rep.h2("Classification set-up")
rep.p(
    "The main task is to predict the label. I compared seven classifiers that cover the usual families: logistic "
    "regression and Gaussian naive Bayes (simple probabilistic models), k-nearest neighbours (instance based), a "
    "single decision tree, two tree ensembles (random forest with 400 trees and XGBoost with 400 boosted trees of "
    "depth 6), and a support vector machine with an RBF kernel. Every model sits in the same pipeline: log transform "
    "of the five skewed columns, iterative imputation, standardisation, then the classifier. The pipeline is fitted on "
    "training rows only, so no statistic from the test rows leaks into preprocessing. The classes are 71 to 29, so "
    "models that support it use balanced class weights.", align="justify")
rep.p(
    "Accuracy alone is misleading with a 71 per cent majority class, since predicting \"patient\" for everyone already "
    "scores 71 per cent. Each model is therefore reported with precision, recall (sensitivity), specificity, F1, "
    "balanced accuracy and ROC-AUC, with a 95 per cent bootstrap interval for AUC from 500 resamples of the test set. "
    "Model choice uses only cross-validated AUC on the training data; the test set is looked at once.", align="justify")
rep.p("Because of what Chapter 2 found, I ran the whole comparison under three splitting protocols:", after=3)
rep.bullets([
    f"Protocol A, the standard split: deduplicated file, stratified random 80/20 split ({n(A['n_train'])} training, "
    f"{n(A['n_test'])} test rows), 5-fold stratified cross-validation.",
    f"Protocol B, grouped by source: every row is assigned to its nearest original patient, and all rows from the same "
    f"source family stay on the same side of the split ({B['n_train']:,} and {B['n_test']:,} rows, "
    f"{R['m7_groups']['n_groups']} families, 5-fold grouped cross-validation).",
    f"Protocol C, real patients only: the {M2['orig_unique']} unique records of the original release, stratified 80/20 "
    f"split ({C['n_train']} and {C['n_test']} rows), cross-validation repeated 5 times over 5 folds because the "
    f"sample is small.",
])
rep.figure(FIG / "f25_nn_distance.png", "Distance from each test record to its nearest training record. Under "
           "protocol A most test rows have a near twin in training.", 14.0)
rep.p(
    f"Figure {F['f25_nn_distance']} shows why protocol A is a problem. The median test row is only {f2(A['nn_dist_median'])} standard "
    f"units from its nearest training row, and that neighbour has the same label {pct(A['nn_label_agreement'], 0)} "
    f"of the time. Under protocol B the median distance rises to {f2(B['nn_dist_median'])}, and among real patients "
    f"(protocol C) it is {f2(C['nn_dist_median'])}.", align="justify")

rep.h2("Protocol A: the standard split")
rows = [mrow(m, v) for m, v in A["models"].items()]
best_a = list(A["models"]).index("Random forest")
rep.table(MHEAD, rows, caption="Protocol A results on the test set. CV AUC is the mean ± SD over five folds of the "
          "training data." + MNOTE, widths=MW, font_size=7.4, highlight_rows=(best_a,))
rep.figure(FIG / "f24_cm_A.png", "Confusion matrices under protocol A. Rows are the actual class, columns the "
           "prediction; each cell shows the count and the share of its row.", 16.4)
am = A["models"]
rep.p(
    f"Random forest and XGBoost both reach {pct(am['Random forest']['acc'])} test accuracy. Random forest makes "
    f"{am['Random forest']['cm'][0][1]} false alarms and misses {am['Random forest']['cm'][1][0]} patients out of "
    f"{n(A['n_test'])}. A McNemar test on their disagreements finds no difference (p = {f2(A['mcnemar']['p'])}). "
    f"The linear models trail far behind at {pct(am['Logistic regression']['acc'], 0)} to "
    f"{pct(am['Gaussian naive Bayes']['acc'], 0)}. If this were the whole story the conclusion would be "
    f"\"use random forest\".", align="justify")

rep.h2("Protocol B: grouped by source patient")
rows = [mrow(m, v) for m, v in B["models"].items()]
rep.table(MHEAD, rows, caption="Protocol B results. Cross-validation folds respect the source families." + MNOTE,
          widths=MW, font_size=7.4, highlight_rows=(list(B["models"]).index("Random forest"),))
rep.figure(FIG / "f24_cm_B.png", "Confusion matrices under protocol B.", 16.4)
bm = B["models"]
rep.p(
    f"Grouping removed the near twins and hurt the models that relied on them: KNN fell from "
    f"{f3(am['k-nearest neighbours (k=15)']['auc'])} to {f3(bm['k-nearest neighbours (k=15)']['auc'])} AUC and the "
    f"SVM from {f3(am['SVM (RBF kernel)']['auc'])} to {f3(bm['SVM (RBF kernel)']['auc'])}. The tree ensembles did not "
    f"move: random forest still scores {pct(bm['Random forest']['acc'])}. So nearness is not what the trees are "
    f"exploiting.", align="justify")
rep.note("Why the trees still win",
         f"A tree can split on a narrow interval around a single value, and in this file single values carry the "
         f"label: a value that only non-patients had in the original data can never appear in a synthetic patient row. "
         f"To check, I wrote a rule with no learning at all. For each column it asks whether the test value was seen "
         f"only among training patients or only among training non-patients, and adds up the votes. That rule scores "
         f"an AUC of {f2(POOL['A']['auc'])} under protocol A and {f2(POOL['B']['auc'])} under protocol B, and "
         f"{pct(POOL['A']['share_with_exclusive_value'], 0)} of test rows contain at least one give-away value. "
         f"On real patients the same rule manages only {f2(POOL['C']['auc'])}.")

rep.h2("Protocol C: real patients")
rows = [mrow(m, v) for m, v in C["models"].items()]
rep.table(MHEAD, rows, caption="Protocol C results on the 114 held-out original patients. CV AUC is the mean ± SD "
          "over 25 folds (5 repeats of 5-fold)." + MNOTE, widths=MW, font_size=7.4,
          highlight_rows=(list(C["models"]).index("SVM (RBF kernel)"), list(C["models"]).index("Logistic regression")))
rep.figure(FIG / "f24_cm_C.png", "Confusion matrices on the 114 held-out original patients (33 non-patients, "
           "81 patients).", 16.4)
cm_ = C["models"]
rep.p(
    f"On real patients the ranking turns upside down. The single decision tree is close to a coin toss (AUC "
    f"{f3(cm_['Decision tree']['auc'])}). Random forest still has a training accuracy of "
    f"{f3(cm_['Random forest']['train_acc'])}, but on the test set it calls almost everyone a patient: it finds "
    f"{cm_['Random forest']['cm'][1][1]} of 81 patients and correctly clears only {cm_['Random forest']['cm'][0][0]} "
    f"of 33 non-patients. Its accuracy of {pct(cm_['Random forest']['acc'], 0)} is below the 71 per cent you get by "
    f"predicting \"patient\" every time. XGBoost and KNN show the same bias towards the majority class.",
    align="justify")
rep.figure(FIG / "f26_roc.png", "ROC curves under the three protocols. The legend gives each model's test AUC.", 16.4)
rep.figure(FIG / "f27_protocol_gap.png", "Test ROC-AUC with 95 per cent bootstrap intervals. Under protocol C the "
           "intervals are wide because the test set holds 114 patients.", 15.0)

rep.h2("Which model outperformed, and why")
sv, lr = cm_["SVM (RBF kernel)"], cm_["Logistic regression"]
rep.p(
    f"On the honest benchmark the RBF support vector machine has the best cross-validated AUC "
    f"({sv['cv_auc_mean']:.3f} ± {sv['cv_auc_sd']:.3f}) and the best test AUC ({sv['auc']:.3f}), with balanced "
    f"accuracy {f3(sv['bacc'])}. Logistic regression is right behind ({lr['cv_auc_mean']:.3f} ± {lr['cv_auc_sd']:.3f} "
    f"in cross-validation, {lr['auc']:.3f} on test). The two models disagree on only "
    f"{C['mcnemar']['table'][0][1] + C['mcnemar']['table'][1][0]} of the 114 test patients and McNemar's test gives "
    f"p = {f2(C['mcnemar']['p'])}, so the data cannot tell them apart. The bootstrap intervals in Figure "
    f"{F['f27_protocol_gap']} overlap "
    f"almost entirely.", align="justify")
rep.p("Three reasons explain why the simple models win on real patients and the ensembles win on the file:", after=3)
rep.bullets([
    "Sample size. 456 training patients is small for a model with thousands of split points. The forest and XGBoost "
    "fit the training data perfectly and carry that noise to the test set; the SVM and logistic regression have far "
    "fewer effective parameters.",
    "Shape of the signal. Chapters 3 and 6 showed that the difference between the classes runs mostly along one "
    "smooth direction (more bilirubin and enzymes, less albumin). A linear boundary on log-scaled values fits that "
    "shape, and the RBF kernel only bends it slightly.",
    "What the ensembles learned on the file. Their 99.7 per cent came from value fingerprints planted by the "
    "augmentation, not from liver physiology. That knowledge has nothing to offer on new patients.",
])
rep.p(
    f"My recommendation is logistic regression. It matches the SVM within noise, produces probabilities a clinician "
    f"can threshold, and its coefficients can be read directly (Figure {F['f28_importance']}). At the default threshold it catches "
    f"{pct(lr['recall'], 0)} of patients while clearing {pct(lr['specificity'], 0)} of non-patients; for screening, "
    f"where a missed patient costs more than a false alarm, the threshold should be lowered to trade specificity "
    f"for sensitivity along the ROC curve.", align="justify")
rep.figure(FIG / "f28_importance.png", "Left: permutation importance for the SVM on the original patients (drop in "
           "AUC when a column is shuffled). Right: logistic regression coefficients per standard deviation.", 16.4)
lc = R["m7"]["lr_coef_C"]
rep.p(
    f"Both views agree on the important inputs: ALP, ALT, direct bilirubin, age and albumin. Albumin has the largest "
    f"negative coefficient ({lc['Albumin']:+.2f} per SD), and age, invisible in the 30k file, matters on real patients "
    f"({lc['Age']:+.2f}). The positive coefficient on total protein next to the negative one on albumin is a "
    f"collinearity effect: holding albumin fixed, more protein means more globulin, which rises in chronic liver "
    f"disease.", align="justify")

rep.h2("Supervised against reinforcement learning")
rep.p(
    "Everything above is supervised learning: each training row comes with the right answer, and the model learns a "
    "mapping from inputs to that answer. Reinforcement learning has no answer key. An agent takes actions in an "
    "environment, receives a reward that may arrive much later, and learns a policy that maximises the total reward "
    "(Sutton and Barto, 2018).", align="justify")
rep.table(["", "Supervised learning", "Reinforcement learning"], [
    ["Feedback", "Correct label for every example", "A scalar reward, often delayed"],
    ["Data", "Fixed dataset collected in advance", "Generated by the agent's own actions"],
    ["Goal", "Predict well on new examples", "Choose actions that maximise long-run reward"],
    ["Main risk", "Overfitting, leakage (Chapter 7)", "Unsafe exploration, reward misspecification"],
    ["Liver example", "Classify a patient from one blood panel", "Adjust a treatment dose visit by visit"],
], caption="Supervised and reinforcement learning compared.", widths=[2.8, 6.8, 6.8], align=["left"] * 3, font_size=8.5)
rep.p(
    "Reinforcement learning is used in game playing, robotics, recommendation systems and, in medicine, for dynamic "
    "treatment regimes, where the action is a dose or a test and the reward is the patient's later outcome. It would "
    "not suit this dataset: there are no actions, no sequence of decisions and no outcome beyond a single label.",
    align="justify")

# =========================================================================== 8
rep.h1("8", "Contemporary issues", "Module 8  /  Contemporary issues")
rep.h2("Synthetic data and leakage")
rep.p(
    "Synthetic and augmented datasets are now common in teaching and in competitions, often without a note saying how "
    "they were made. This project shows the cost. The 30k file is easy to mistake for a large real cohort, it "
    "rewards the most flexible models with near-perfect scores, and it hides two real effects (age and gender). "
    "Kaufman and colleagues (2012) call this leakage: information in the training data that would not be available "
    "when the model is used. The checks that caught it were simple and are worth doing on any dataset: count "
    "duplicates, look for the original source, measure how close test rows are to training rows, and distrust any "
    "accuracy above what the exploratory analysis says is plausible. Chapters 3 and 6 showed heavy overlap between "
    "the classes, and that overlap, not the 99.7 per cent, is what protocol C measured.", align="justify")
rep.h2("Fairness")
fr = R["m8_fairness"]
fc = fr["C"]
rep.p(
    f"I checked whether the recommended kind of model works equally well for men and women. On the original test "
    f"patients the SVM's sensitivity was {pct(fc['Male']['sensitivity'], 0)} for men and "
    f"{pct(fc['Female']['sensitivity'], 0)} for women, with specificity {pct(fc['Male']['specificity'], 0)} and "
    f"{pct(fc['Female']['specificity'], 0)}. The test set has only {fc['Female']['n']} women, so the gap is uncertain, "
    f"but it points the same way as the original data: women in this sample are fewer and less often patients, and "
    f"the model finds their disease less reliably. A screening tool built on this data would need a check on a larger "
    f"and more balanced female sample before use. On the 30k file the same check shows no gap at all, because the "
    f"augmentation removed gender from the problem.", align="justify")
rep.h2("Reproducibility")
rep.p(
    f"All random steps use seed {R['env']['seed']}. Imputation and scaling are fitted inside the model pipeline, on "
    f"training rows only. The analysis script writes every number quoted here to results/results.json, and the report "
    f"is generated from that file, so the text and the results cannot drift apart. The original UCI file used for "
    f"the provenance check is stored with the project.", align="justify")

# =========================================================================== 9
rep.h1("9", "Conclusions", "Summary")
rep.p("Going back to the three questions from Chapter 1:", after=3)
rep.bullets([
    f"The file is not clean in the way it looks. Over a third of it is exact duplicates, and the rest was generated "
    f"from {M2['orig_unique']} real patients by mixing their values within each class. That process leaks the label "
    f"and erases the age and gender effects. Missing values are few and are not completely random; EM imputation "
    f"handles them well.",
    "Bilirubin, the transaminases and alkaline phosphatase separate patients from non-patients, with albumin pointing "
    "the other way. Factor analysis recovers the three clinical groups of the liver panel (synthesis, injury, "
    "excretion), and clustering finds groups that look like stages of damage, though none of them matches the label.",
    f"On real, unseen patients the best achievable ROC-AUC is about 0.73 in cross-validation. An RBF SVM and logistic "
    f"regression share first place; I recommend logistic regression for its transparency. The 99.7 per cent accuracy "
    f"of random forest on the course file measures the data generator, not the disease.",
], numbered=True)
rep.p(
    "The limitations are real. Protocol C has only 114 test patients, so its intervals are wide. The label in the "
    "original data comes from a single hospital system and its definition of a \"liver patient\" is not documented. "
    "And with no follow-up dates, nothing here says anything about how the disease progresses. A larger, dated cohort "
    "would allow the time-series analysis that Module 3 describes and a far tighter estimate of model performance.",
    align="justify")

# =========================================================================== references
rep.h1("", "References")
refs = [
    "Aggarwal, C. C. (2015). _Data Mining: The Textbook_. Springer.",
    "Chawla, N. V., Bowyer, K. W., Hall, L. O., and Kegelmeyer, W. P. (2002). SMOTE: synthetic minority over-sampling technique. _Journal of Artificial Intelligence Research_, 16, 321 to 357.",
    "Chen, T., and Guestrin, C. (2016). XGBoost: a scalable tree boosting system. _Proceedings of the 22nd ACM SIGKDD Conference_, 785 to 794.",
    "Dempster, A. P., Laird, N. M., and Rubin, D. B. (1977). Maximum likelihood from incomplete data via the EM algorithm. _Journal of the Royal Statistical Society, Series B_, 39(1), 1 to 38.",
    "Enders, C. K. (2010). _Applied Missing Data Analysis_. The Guilford Press.",
    "Fraley, C. (1998). Algorithms for model-based Gaussian hierarchical clustering. _SIAM Journal on Scientific Computing_, 20(1), 270 to 281.",
    "Jambu, M. (1991). _Exploratory and Multivariate Data Analysis_. Academic Press.",
    "Kaufman, S., Rosset, S., Perlich, C., and Stitelman, O. (2012). Leakage in data mining: formulation, detection, and avoidance. _ACM Transactions on Knowledge Discovery from Data_, 6(4), article 15.",
    "Kohonen, T. (1982). Self-organized formation of topologically correct feature maps. _Biological Cybernetics_, 43, 59 to 69.",
    "Levina, E., and Bickel, P. J. (2004). Maximum likelihood estimation of intrinsic dimension. _Advances in Neural Information Processing Systems_, 17.",
    "Martinez, W. L., Martinez, A. R., and Solka, J. L. (2017). _Exploratory Data Analysis with MATLAB_ (3rd ed.). CRC Press.",
    "McNemar, Q. (1947). Note on the sampling error of the difference between correlated proportions or percentages. _Psychometrika_, 12(2), 153 to 157.",
    "Mukhiya, S. K., and Ahmed, U. (2020). _Hands-On Exploratory Data Analysis with Python_. Packt Publishing.",
    "Pedregosa, F., et al. (2011). Scikit-learn: machine learning in Python. _Journal of Machine Learning Research_, 12, 2825 to 2830.",
    "Ramana, B. V., Babu, M. S. P., and Venkateswarlu, N. B. (2012). ILPD (Indian Liver Patient Dataset). UCI Machine Learning Repository, https://archive.ics.uci.edu/dataset/225/ilpd+indian+liver+patient+dataset",
    "Rubin, D. B. (1976). Inference and missing data. _Biometrika_, 63(3), 581 to 592.",
    "Sutton, R. S., and Barto, A. G. (2018). _Reinforcement Learning: An Introduction_ (2nd ed.). MIT Press.",
    "Tenenbaum, J. B., de Silva, V., and Langford, J. C. (2000). A global geometric framework for nonlinear dimensionality reduction. _Science_, 290, 2319 to 2323.",
    "van der Maaten, L., and Hinton, G. (2008). Visualizing data using t-SNE. _Journal of Machine Learning Research_, 9, 2579 to 2605.",
    "von Luxburg, U. (2007). A tutorial on spectral clustering. _Statistics and Computing_, 17, 395 to 416.",
    "Zahn, C. T. (1971). Graph-theoretical methods for detecting and describing gestalt clusters. _IEEE Transactions on Computers_, C-20(1), 68 to 86.",
]
for ref in refs:
    para = rep.p(ref, size=9, after=4)
    para.paragraph_format.left_indent = Cm(0.6)
    para.paragraph_format.first_line_indent = Cm(-0.6)

# =========================================================================== appendix
rep.h1("A", "Appendix: running the project")
rep.p("The project folder is laid out as follows.", after=3)
rep.table(["Path", "Contents"], [
    ["data/liver.csv", "The course file, as downloaded"],
    ["data/ilpd_original_583.csv", "The original UCI release, used for the provenance check and protocol C"],
    ["src/analysis.py", "Every analysis step, in syllabus order; writes figures and results"],
    ["src/style.py", "Shared chart style"],
    ["src/build_report.py, src/docx_theme.py", "Generates this report from results/results.json"],
    ["figures/", f"The {rep.fig_n} figures in this report, 220 dpi PNG"],
    ["results/results.json, results/tables/", "Every quoted number, and the tables as CSV"],
], caption="Project files.", widths=[6.0, 10.4], align=["left", "left"], font_size=8.5)
rep.p("To reproduce everything from a fresh copy of the folder:", after=3)
for line in ("pip install pandas numpy scipy scikit-learn statsmodels xgboost matplotlib python-docx",
             "python src/analysis.py", "python src/build_report.py"):
    para = rep.p(line, size=8.5, font="Consolas", after=1)
    para.paragraph_format.left_indent = Cm(0.6)
rep.p("")
rep.p(f"Software versions used: Python {R['env']['python']}, pandas {R['env']['pandas']}, NumPy {R['env']['numpy']}, "
      f"scikit-learn {R['env']['sklearn']}, statsmodels {R['env']['statsmodels']}, XGBoost {R['env']['xgboost']}.",
      size=9, color=MUTED)

assert rep.fig_order == ORDER, (rep.fig_order, ORDER)
rep.save(OUT / "Liver_EDA_Report.docx")
print("figures:", rep.fig_n, "tables:", rep.tab_n)
