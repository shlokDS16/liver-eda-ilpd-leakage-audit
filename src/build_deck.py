"""Builds report/Liver_EDA_Presentation.pptx: a 7-slide deck for a ~3 minute recorded talk.
The narration for each slide is stored in that slide's speaker notes.

Run after src/analysis.py:  python src/build_deck.py
"""
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "figures"
R = json.load(open(ROOT / "results" / "results.json"))
A, C = R["m7"]["A"], R["m7"]["C"]

BG = RGBColor(0xFB, 0xFB, 0xFA)
INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x6B, 0x72, 0x80)
RULE = RGBColor(0xD9, 0xD4, 0xCC)
OX = RGBColor(0x8C, 0x2F, 0x39)
HEAD = "Bahnschrift SemiBold"
BODY = "Bahnschrift"
LIGHT = "Bahnschrift Light"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height
MX = Inches(0.75)   # side margin


def text(slide, x, y, w, h, runs, size=18, color=INK, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=1.1):
    """runs: str, or list of paragraphs; a paragraph is a str or a list of (text, overrides) tuples."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        parts = [(para, {})] if isinstance(para, str) else para
        for t, o in parts:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.color.rgb = o.get("color", color)
            f.bold = o.get("bold", False)
        if isinstance(para, list) and para and "after" in para[0][1]:
            p.space_after = Pt(para[0][1]["after"])
    return tb


def rule(slide, x, y, w, h=Emu(9525), color=RULE):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def picture(slide, path, x, y, w, h):
    """Fit the image inside the box, keeping its aspect ratio, aligned top-left."""
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    return slide.shapes.add_picture(str(path), x, y, int(iw * scale), int(ih * scale))


def new_slide(n, notes):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    if n > 1:
        text(s, MX, H - Inches(0.55), Inches(6), Inches(0.3), "BCSE331L  Exploratory Data Analysis", size=10, color=MUTED)
        text(s, W - MX - Inches(1), H - Inches(0.55), Inches(1), Inches(0.3), f"{n} / 7", size=10, color=MUTED,
             align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


def headline(s, t, y=Inches(0.65), w=Inches(11.8), size=34):
    text(s, MX, y, w, Inches(1.1), t, size=size, font=HEAD, spacing=1.0)


def stat(s, x, y, w, big, small, color=INK, big_size=40):
    text(s, x, y, w, Inches(0.8), big, size=big_size, font=LIGHT, color=color, spacing=1.0)
    text(s, x, y + Inches(0.78), w, Inches(0.6), small, size=13, color=MUTED, spacing=1.05)


# ------------------------------------------------------------------ narration
NOTES = [
    # 1
    "Hi, I'm Shlok Kumar Goenka, register number 23BDS0316, and this is my Exploratory Data Analysis course "
    "project under Balaji G N. I analysed a liver patient dataset with about thirty thousand rows. The title gives "
    "away the twist: those rows come from only five hundred and seventy real patients.",
    # 2
    "Each row is a routine liver blood panel: age, gender, two bilirubin values, three liver enzymes, total "
    "protein, albumin and the albumin to globulin ratio, plus a label saying whether the person is a liver "
    "patient. I set three questions. Is the data clean? Which measurements separate patients from everyone else? "
    "And how well can a model predict the label for someone it has never seen?",
    # 3
    "Cleaning gave the biggest surprise. Eleven thousand three hundred rows were exact duplicates. The first rows "
    "matched the original 583-patient dataset from the UCI repository, so I downloaded it and compared. Every lab "
    "value in the added rows came from an original patient of the same class. That means the values themselves "
    "leak the label. It also erased a real effect: in the original data men have a higher disease rate than "
    "women, but in this file both sit at 71 per cent. For the missing values I used the EM algorithm, which cut "
    "the error by more than a third compared with mean imputation.",
    # 4
    "The exploratory analysis told a consistent story. Bilirubin, the liver enzymes and alkaline phosphatase are "
    "higher in patients, and albumin is lower. Factor analysis recovered three groups that textbooks describe: "
    "protein synthesis, cell injury, and bilirubin handling. Clustering found four groups that look like stages of "
    "damage, but none of them matches the label, because patients drift away from normal in many directions.",
    # 5
    "Then the models. I split the data 80 to 20 and compared seven classifiers under three protocols. On a random "
    "split, random forest and XGBoost reach 99.7 per cent accuracy. That looks perfect. But a simple lookup rule "
    "with no learning at all reaches an AUC of 0.94 on the same split, so the forest is reading the generator's "
    "fingerprints. When I trained and tested only on the 570 real patients, the forest fell apart. It cleared "
    "just 4 of 33 healthy people.",
    # 6
    "On real patients, the RBF support vector machine and logistic regression tie, with a cross-validated AUC of "
    "about 0.73, and a McNemar test cannot separate them. I recommend logistic regression. It is just as accurate, "
    "it gives probabilities a doctor can threshold, and its coefficients can be read directly. Alkaline "
    "phosphatase, ALT, direct bilirubin, age and albumin carry most of the weight.",
    # 7
    "So my main lesson is simple. Check where your data came from before you trust a score. Here, ninety-nine "
    "per cent accuracy measured the data generator, not the disease. The full report and code are in my GitHub "
    "repository. Thank you.",
]

# ------------------------------------------------------------------ 1  title
s = new_slide(1, NOTES[0])
text(s, MX, Inches(0.8), Inches(7), Inches(0.3), "BCSE331L  EXPLORATORY DATA ANALYSIS  /  COURSE PROJECT",
     size=11, color=OX)
text(s, MX, Inches(1.9), Inches(7.6), Inches(2.6),
     [[("Thirty thousand rows,", {})], [("five hundred and seventy patients", {"color": OX})]],
     size=40, font=HEAD, spacing=0.95)
text(s, MX, Inches(4.35), Inches(6.8), Inches(0.9),
     "An exploratory analysis of liver function records, and why a 99.7 per cent model learned almost nothing "
     "about the liver.", size=17, color=MUTED, spacing=1.15)
text(s, MX, Inches(6.05), Inches(7), Inches(0.9),
     [[("Shlok Kumar Goenka", {"font": HEAD, "size": 15, "color": INK}), ("   23BDS0316", {"size": 15})],
      [("Faculty: Balaji G N", {"size": 13})]], size=13, color=MUTED, spacing=1.2)
rule(s, Inches(9.05), Inches(1.9), Emu(12700), Inches(4.6), RULE)
stat(s, Inches(9.45), Inches(1.9), Inches(3.2), f"{R['m1']['n_rows']:,}", "rows in the course file")
stat(s, Inches(9.45), Inches(3.5), Inches(3.2), f"{R['m2']['orig_unique']}", "real patients behind them")
stat(s, Inches(9.45), Inches(5.1), Inches(3.2), f"{C['models']['SVM (RBF kernel)']['cv_auc_mean']:.2f}",
     "honest ROC-AUC on unseen patients", color=OX)

# ------------------------------------------------------------------ 2  data and questions
s = new_slide(2, NOTES[1])
headline(s, "A routine liver panel, 30,691 times")
groups = [("Bilirubin", "total, direct"), ("Enzymes", "ALP, ALT (SGPT), AST (SGOT)"),
          ("Proteins", "total protein, albumin, A/G ratio"), ("Person", "age, gender"),
          ("Label", "liver patient or not (71% patients)")]
y = Inches(2.0)
for name, val in groups:
    text(s, MX, y, Inches(1.7), Inches(0.5), name, size=16, font=HEAD)
    text(s, MX + Inches(1.8), y, Inches(4.2), Inches(0.5), val, size=16, color=MUTED)
    y += Inches(0.68)
rule(s, MX, Inches(1.75), Inches(5.9))
qs = ["Is the data clean enough to trust?",
      "Which measurements separate patients from everyone else?",
      "How well can a model predict the label for someone new?"]
y = Inches(2.0)
for i, q in enumerate(qs, 1):
    text(s, Inches(7.4), y, Inches(0.6), Inches(0.8), f"{i}", size=40, font=LIGHT, color=OX, spacing=1.0)
    text(s, Inches(8.1), y + Inches(0.12), Inches(4.5), Inches(0.9), q, size=19, spacing=1.1)
    y += Inches(1.35)

# ------------------------------------------------------------------ 3  provenance
s = new_slide(3, NOTES[2])
headline(s, "The extra rows were copied from real patients")
stat(s, MX, Inches(1.9), Inches(3.4), f"{R['m2']['n_exact_dup']:,}", "exact duplicate rows removed")
stat(s, MX, Inches(3.3), Inches(3.4), f"{R['m2']['n_orig_rows_found']}", "rows identical to the 583-patient UCI release")
stat(s, MX, Inches(4.7), Inches(3.4), "100%", "of lab values in added rows come from same-class patients", color=OX)
picture(s, FIG / "f29_provenance.png", Inches(4.6), Inches(1.95), Inches(8.0), Inches(4.4))

# ------------------------------------------------------------------ 4  EDA
s = new_slide(4, NOTES[3])
headline(s, "Bilirubin and enzymes rise, albumin falls")
picture(s, FIG / "f12_dotchart.png", MX, Inches(1.85), Inches(7.0), Inches(4.6))
x = Inches(8.4)
rule(s, x, Inches(1.95), Inches(4.2))
text(s, x, Inches(2.15), Inches(4.2), Inches(1.6),
     [[("Three factors", {"font": HEAD, "size": 19})],
      [("protein synthesis, cell injury, bilirubin handling", {"color": MUTED, "size": 15})]], spacing=1.15)
rule(s, x, Inches(3.55), Inches(4.2))
text(s, x, Inches(3.75), Inches(4.2), Inches(1.6),
     [[("Four clusters", {"font": HEAD, "size": 19})],
      [("look like stages of damage, but none matches the label", {"color": MUTED, "size": 15})]], spacing=1.15)
rule(s, x, Inches(5.15), Inches(4.2))
text(s, x, Inches(5.35), Inches(4.2), Inches(1.2),
     [[("About 4 real dimensions", {"font": HEAD, "size": 19})],
      [("out of 9 measured columns", {"color": MUTED, "size": 15})]], spacing=1.15)

# ------------------------------------------------------------------ 5  evaluation
s = new_slide(5, NOTES[4])
headline(s, "99.7% on the file, 0.73 on real patients")
rows = [("A", "Random 80/20 split of the file", f"{A['models']['Random forest']['acc'] * 100:.1f}% accuracy"),
        ("B", "Grouped by source patient", "trees still near 99%"),
        ("C", "Only the 570 real patients", f"best AUC {C['models']['SVM (RBF kernel)']['cv_auc_mean']:.2f}")]
y = Inches(1.95)
for k, desc, res in rows:
    text(s, MX, y, Inches(0.5), Inches(0.6), k, size=30, font=LIGHT, color=OX, spacing=1.0)
    text(s, MX + Inches(0.6), y + Inches(0.02), Inches(4.4), Inches(0.9),
         [[(desc, {"size": 16})], [(res, {"size": 14, "color": MUTED})]], spacing=1.1)
    y += Inches(1.05)
rule(s, MX, Inches(5.2), Inches(4.9))
text(s, MX, Inches(5.35), Inches(4.9), Inches(0.9),
     [[("A lookup rule with no learning scores AUC ", {"size": 15}),
       (f"{R['m7_pool_rule']['A']['auc']:.2f}", {"size": 15, "color": OX, "font": HEAD}),
       (" on split A", {"size": 15})]], spacing=1.1)
picture(s, FIG / "f27_protocol_gap.png", Inches(6.0), Inches(1.85), Inches(6.6), Inches(4.7))

# ------------------------------------------------------------------ 6  winner
s = new_slide(6, NOTES[5])
headline(s, "Logistic regression: as good, easier to explain")
sv, lr, rf = C["models"]["SVM (RBF kernel)"], C["models"]["Logistic regression"], C["models"]["Random forest"]
stat(s, MX, Inches(1.95), Inches(2.4), f"{sv['cv_auc_mean']:.3f}", "SVM (RBF), cross-validated AUC", big_size=34)
stat(s, MX + Inches(2.6), Inches(1.95), Inches(2.4), f"{lr['cv_auc_mean']:.3f}", "logistic regression", big_size=34,
     color=OX)
stat(s, MX, Inches(3.55), Inches(2.4), f"p = {C['mcnemar']['p']:.2f}", "McNemar test, no difference", big_size=34)
stat(s, MX + Inches(2.6), Inches(3.55), Inches(2.4), f"{rf['cm'][0][0]} of {sum(rf['cm'][0])}",
     "healthy people cleared by random forest", big_size=34)
text(s, MX, Inches(5.25), Inches(5.0), Inches(1.0),
     "Simple models win on 456 training patients; the forest memorised the generator.", size=15, color=MUTED,
     spacing=1.15)
picture(s, FIG / "f28_importance.png", Inches(6.15), Inches(2.0), Inches(6.5), Inches(4.4))

# ------------------------------------------------------------------ 7  takeaway
s = new_slide(7, NOTES[6])
text(s, MX, Inches(1.6), Inches(9.5), Inches(2.4),
     [[("Check where the data came from", {})], [("before trusting the score.", {"color": OX})]],
     size=44, font=HEAD, spacing=1.0)
rule(s, MX, Inches(3.75), Inches(11.8))
takeaways = [("Cleaning", "found 36.9% duplicates and a class-conditional generator"),
             ("Exploration", "bilirubin, enzymes and albumin carry the signal"),
             ("Evaluation", "honest AUC about 0.73; logistic regression recommended")]
x = MX
for name, val in takeaways:
    text(s, x, Inches(4.05), Inches(3.7), Inches(1.4),
         [[(name, {"font": HEAD, "size": 18})], [(val, {"size": 15, "color": MUTED})]], spacing=1.15)
    x += Inches(4.0)
text(s, MX, Inches(5.9), Inches(8), Inches(0.4), "Thank you", size=16, color=MUTED)

out = ROOT / "report" / "Liver_EDA_Presentation.pptx"
prs.save(out)
(ROOT / "report" / "video_script.txt").write_text(
    "\n\n".join(f"Slide {i}\n{t}" for i, t in enumerate(NOTES, 1)) + "\n", encoding="utf-8")
words = sum(len(t.split()) for t in NOTES)
print(f"saved {out.name}; script {words} words, about {words / 150:.1f} min at 150 wpm")
