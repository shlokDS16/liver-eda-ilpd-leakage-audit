"""Word theme for the report: Georgia body, Bahnschrift display type, oxblood accent,
booktabs-style tables. Kept separate from the content so the look can change in one place."""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x6B, 0x72, 0x80)
OX = RGBColor(0x8C, 0x2F, 0x39)
TEAL = RGBColor(0x2F, 0x5D, 0x62)
BODY = "Georgia"
DISPLAY = "Bahnschrift"
TINT = "F4EFE8"


def _font(run_or_style, name, size=None, color=None, bold=None, italic=None):
    f = run_or_style.font
    f.name = name
    rpr = run_or_style.element.get_or_add_rPr() if hasattr(run_or_style.element, "get_or_add_rPr") else None
    if rpr is not None:
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(a), name)
    if size:
        f.size = Pt(size)
    if color is not None:
        f.color.rgb = color
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic


class Report:
    def __init__(self):
        self.doc = Document()
        self.fig_n = 0
        self.fig_order = []
        self.tab_n = 0
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.3)
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
        sec.header_distance = sec.footer_distance = Cm(1.0)
        st = self.doc.styles
        _font(st["Normal"], BODY, 10, INK)
        pf = st["Normal"].paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.22
        for name, size, before, after, color in (("Heading 1", 20, 0, 10, INK), ("Heading 2", 12.5, 14, 4, OX),
                                                 ("Heading 3", 10.5, 10, 2, INK)):
            s = st[name]
            _font(s, DISPLAY, size, color, bold=False, italic=False)
            s.paragraph_format.space_before = Pt(before)
            s.paragraph_format.space_after = Pt(after)
            s.paragraph_format.keep_with_next = True
            s.paragraph_format.line_spacing = 1.05
        _font(st["Caption"], DISPLAY, 8.5, MUTED, bold=False, italic=False)
        st["Caption"].paragraph_format.space_after = Pt(10)
        st["Caption"].paragraph_format.space_before = Pt(3)
        for lvl in (1, 2):
            try:
                s = st[f"TOC {lvl}"]
            except KeyError:
                s = st.add_style(f"TOC {lvl}", 1)
            _font(s, DISPLAY if lvl == 1 else BODY, 10 if lvl == 1 else 9.5, INK)
            s.paragraph_format.space_after = Pt(2 if lvl == 2 else 4)
            s.paragraph_format.space_before = Pt(6 if lvl == 1 else 0)
            s.paragraph_format.left_indent = Cm(0 if lvl == 1 else 0.8)
        self._lists()

    # ------------------------------------------------------------------ basics
    def _lists(self):
        for name in ("List Bullet", "List Number"):
            s = self.doc.styles[name]
            _font(s, BODY, 10, INK)
            s.paragraph_format.space_after = Pt(3)

    def p(self, text="", size=None, color=None, italic=False, align=None, after=None, before=None, keep=False, font=None):
        para = self.doc.add_paragraph()
        self._runs(para, text, size=size, color=color, italic=italic, font=font)
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == "justify":
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        if before is not None:
            para.paragraph_format.space_before = Pt(before)
        if keep:
            para.paragraph_format.keep_with_next = True
        return para

    def _runs(self, para, text, size=None, color=None, italic=False, font=None):
        """Very small markup: **bold** and _italic_ segments."""
        import re
        parts = re.split(r"(\*\*[^*]+\*\*|(?<![A-Za-z0-9])_[^_]+_(?![A-Za-z0-9]))", text)
        for part in parts:
            if not part:
                continue
            bold = part.startswith("**")
            it = part.startswith("_") and part.endswith("_") and len(part) > 2
            t = part[2:-2] if bold else (part[1:-1] if it else part)
            r = para.add_run(t)
            _font(r, font or BODY, size, color, bold=bold or None, italic=(italic or it) or None)
        return para

    def bullets(self, items, numbered=False):
        for i, it in enumerate(items, 1):
            if numbered:
                para = self.doc.add_paragraph()
                r = para.add_run(f"{i}.\t")
                _font(r, DISPLAY, 10, OX)
                para.paragraph_format.left_indent = Cm(0.75)
                para.paragraph_format.first_line_indent = Cm(-0.75)
                para.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
                para.paragraph_format.space_after = Pt(3)
            else:
                para = self.doc.add_paragraph(style="List Bullet")
            self._runs(para, it)

    def h1(self, number, title, kicker=None):
        self.page_break()
        if kicker:
            k = self.doc.add_paragraph()
            r = k.add_run(kicker.upper())
            _font(r, DISPLAY, 8, OX)
            r.font.spacing = Pt(1.2) if hasattr(r.font, "spacing") else None
            self._tracking(r, 30)
            k.paragraph_format.space_after = Pt(2)
        h = self.doc.add_paragraph(style="Heading 1")
        if number:
            rn = h.add_run(f"{number}  ")
            _font(rn, DISPLAY, 20, OX)
        rt = h.add_run(title)
        _font(rt, DISPLAY, 20, INK)
        self._rule(h, OX, 8, space=6)
        return h

    def h2(self, title):
        return self.doc.add_paragraph(title, style="Heading 2")

    def h3(self, title):
        return self.doc.add_paragraph(title, style="Heading 3")

    def page_break(self):
        if len(self.doc.paragraphs) == 0:
            return
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    @staticmethod
    def _tracking(run, val):
        rpr = run._r.get_or_add_rPr()
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(val))
        rpr.append(sp)

    @staticmethod
    def _rule(para, color, size=6, space=4, side="bottom"):
        ppr = para._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size))
        b.set(qn("w:space"), str(space))
        b.set(qn("w:color"), str(color))
        bdr.append(b)
        ppr.append(bdr)

    @staticmethod
    def _shade(para_or_cell, fill):
        el = para_or_cell._p.get_or_add_pPr() if hasattr(para_or_cell, "_p") else para_or_cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        el.append(shd)

    def note(self, label, text):
        """Callout: tinted paragraph with an oxblood bar on the left."""
        para = self.doc.add_paragraph()
        r = para.add_run(label.upper() + "   ")
        _font(r, DISPLAY, 8, OX)
        self._tracking(r, 20)
        self._runs(para, text, size=9.5)
        ppr = para._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:left")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "24")
        b.set(qn("w:space"), "8")
        b.set(qn("w:color"), "8C2F39")
        bdr.append(b)
        ppr.append(bdr)
        self._shade(para, TINT)
        para.paragraph_format.left_indent = Cm(0.25)
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after = Pt(10)
        para.paragraph_format.line_spacing = 1.2
        return para

    # ------------------------------------------------------------------ figures and tables
    def figure(self, path, caption, width_cm=16.4):
        self.fig_n += 1
        self.fig_order.append(Path(path).stem)
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.space_after = Pt(0)
        para.add_run().add_picture(str(path), width=Cm(width_cm))
        cap = self.doc.add_paragraph(style="Caption")
        r = cap.add_run(f"Figure {self.fig_n}   ")
        _font(r, DISPLAY, 8.5, OX)
        self._runs(cap, caption, size=8.5, color=MUTED, font=DISPLAY)
        return self.fig_n

    def table(self, header, rows, caption=None, widths=None, align=None, font_size=8.5, highlight_rows=()):
        if caption:
            self.tab_n += 1
            cap = self.doc.add_paragraph(style="Caption")
            cap.paragraph_format.keep_with_next = True
            cap.paragraph_format.space_after = Pt(3)
            r = cap.add_run(f"Table {self.tab_n}   ")
            _font(r, DISPLAY, 8.5, OX)
            self._runs(cap, caption, size=8.5, color=MUTED, font=DISPLAY)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for side, sz, col in (("top", 12, "1F2328"), ("bottom", 12, "1F2328"), ("insideH", 2, "E4DED5"),
                              ("left", 0, "FFFFFF"), ("right", 0, "FFFFFF"), ("insideV", 0, "FFFFFF")):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single" if sz else "nil")
            b.set(qn("w:sz"), str(sz))
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), col)
            borders.append(b)
        tblPr.append(borders)
        mar = OxmlElement("w:tblCellMar")
        for side, v in (("top", 30), ("bottom", 30), ("left", 80), ("right", 80)):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), str(v))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tblPr.append(mar)
        align = align or (["left"] + ["right"] * (len(header) - 1))
        for ri, row in enumerate([header] + list(rows)):
            cells = t.rows[ri].cells
            for ci, val in enumerate(row):
                cell = cells[ci]
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                para.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                                  "center": WD_ALIGN_PARAGRAPH.CENTER}[align[ci]]
                if ri == 0:
                    r = para.add_run(str(val))
                    _font(r, DISPLAY, font_size, INK)
                    self._shade(cell, TINT)
                else:
                    self._runs(para, str(val), size=font_size, font=DISPLAY if ci == 0 else BODY)
                    if (ri - 1) in highlight_rows:
                        self._shade(cell, "F7E9E7")
            if ri == 0:
                trPr = t.rows[0]._tr.get_or_add_trPr()
                th = OxmlElement("w:tblHeader")
                th.set(qn("w:val"), "true")
                trPr.append(th)
                b = OxmlElement("w:cantSplit")
                trPr.append(b)
        if len(rows) <= 15:   # keep short tables on one page
            for row in t.rows[:-1]:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.keep_with_next = True
        if widths:
            for ci, w in enumerate(widths):
                for row in t.rows:
                    row.cells[ci].width = Cm(w)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        return t

    # ------------------------------------------------------------------ fields, header, footer
    @staticmethod
    def _field(para, instr, placeholder=""):
        r = para.add_run()
        b = OxmlElement("w:fldChar")
        b.set(qn("w:fldCharType"), "begin")
        r._r.append(b)
        r2 = para.add_run()
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        r2._r.append(it)
        r3 = para.add_run()
        s = OxmlElement("w:fldChar")
        s.set(qn("w:fldCharType"), "separate")
        r3._r.append(s)
        r4 = para.add_run(placeholder)
        r5 = para.add_run()
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), "end")
        r5._r.append(e)
        return [r, r2, r3, r4, r5]

    def toc(self):
        para = self.doc.add_paragraph()
        self._field(para, 'TOC \\o "1-2" \\h \\z \\u', "Right-click and choose Update Field to build the contents.")

    def new_body_section(self, running_title):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = hp.add_run(running_title)
        _font(r, DISPLAY, 7.5, MUTED)
        self._tracking(r, 10)
        self._rule(hp, "D9D4CC", 4, 4)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for rr in self._field(fp, "PAGE", "1"):
            _font(rr, DISPLAY, 8, OX)
        # restart numbering at 1
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:start"), "1")
        sec._sectPr.append(pg)
        return sec

    def save(self, path):
        self.doc.save(str(path))
